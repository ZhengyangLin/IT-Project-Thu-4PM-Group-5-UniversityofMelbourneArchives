"""Service implementation unit tests against temporary folders and fake workers.

No PaddleOCR inference, process pool, or production status-file writes.
"""
from __future__ import annotations
import asyncio
import json
from concurrent.futures.process import BrokenProcessPool
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
import pytest
from openpyxl import load_workbook
from backend_test_helpers import module

@pytest.fixture
def env(monkeypatch,tmp_path):
    m=module('services.impl.Unimelb_service_impl')
    root=tmp_path/'plan_sample';root.mkdir()
    out=tmp_path/'output'
    monkeypatch.setattr(m,'PROJECT_DIR',tmp_path)
    monkeypatch.setattr(m,'PLAN_SAMPLE_DIR',root)
    monkeypatch.setattr(m,'UNIMELB_OCR_OUTPUT_ROOT',out)
    monkeypatch.setattr(m,'UNIMELB_CONFIG_PATH',tmp_path/'config.yaml')
    monkeypatch.setattr(m,'IMAGE_STATUS_FILE',out/'status.json')
    return SimpleNamespace(m=m,root=root,out=out,tmp=tmp_path)

@pytest.fixture
def svc(env):
    return env.m.UnimelbServiceImpl(status_file=env.out/'test-status.json')

@pytest.fixture
def sample(env):
    folder=env.root/'Sample'; node=folder/'N001';node.mkdir(parents=True)
    path=node/'a.png';path.write_bytes(b'example')
    return path

def state(env,path,status=0,**kwargs):
    result=env.m.UnimelbServiceImpl._new_image_state(path,env.m.UNIMELB_SHARED_TASK_UUID,status,100)
    result.update(kwargs)
    return result

@pytest.mark.parametrize('value,expected',[
    (0,0),(1,1),(2,2),(3,3),('0',0),('1',1),('2',2),('3',3),
    ('queued',0),('not_started',0),('running',1),('success',2),('failed',3),
    ('invalid',0),(None,0),(-1,0),(44,0),('',0)])
def test_status_normalization(env,value,expected):
    assert env.m.UnimelbServiceImpl._normalize_image_status(value)==expected

@pytest.mark.parametrize('value,expected',[(0,0),(1,1),('0',0),('1',1),(' 1 ',1),(' 0 ',0),('yes',0),(None,0),(-1,0),(2,0)])
def test_review_normalization(env,value,expected):
    assert env.m.UnimelbServiceImpl._normalize_manual_reviewed(value)==expected


def test_image_listing_filters_and_sorting(svc,env,sample):
    (env.root/'bates_smart_mccutcheon').mkdir()
    (env.root/'bates_smart_mccutcheon'/'ignore.png').write_bytes(b'x')
    (sample.parent/'b.JPG').write_bytes(b'y')
    (sample.parent/'ignore.txt').write_text('no')
    assert [f.name for f in svc._image_folders()]==['Sample']
    assert [f.name for f in svc._collect_images()]==['a.png','b.JPG']
    assert svc._image_key(sample)=='N001/a.png'
    assert svc._image_sort_key(sample)==('Sample','N001/a.png')
    assert 'Sample' in svc._display_image_path(sample)
    assert svc._item_output_dir(Path('/out'),'N','a.png')==Path('/out/items/N/a.png')
    assert svc._now_ms()>0


def test_missing_root(svc,env):
    env.root.rmdir()
    assert svc._image_folders()==[] and svc._collect_images()==[]
    with pytest.raises(FileNotFoundError): svc._validate_folders()

@pytest.mark.parametrize('folder,node,name,found',[
    ('sample','N001','a.png',True),('SAMPLE','N001','a.png',True),
    ('missing','N001','a.png',False),('Sample','N001','bad.png',False),
    ('Sample','../bad','a.png',False),('Sample','N001','../bad',False),
    ('Sample','N001','notimage.txt',False)])
def test_resolve_image(svc,env,sample,folder,node,name,found):
    (sample.parent/'notimage.txt').write_text('no')
    result=svc.resolve_image(folder,node,name)
    assert (result is not None) is found
    if found: assert result==sample.resolve()


def test_new_state_and_save_roundtrip(svc,env,sample):
    s=svc._new_image_state(sample,'TASK',1,500)
    assert s['status']==1 and s['manual_reviewed']==0 and s['image_name']=='a.png'
    svc._image_statuses['N001/a.png']=s
    svc._persist_image_statuses_locked()
    saved=json.loads(svc._status_file.read_text(encoding='utf-8'))
    assert saved['N001/a.png']['updated_at']==500
    loaded,migrated=svc._load_image_statuses()
    assert loaded['N001/a.png']['status']==1
    assert not migrated


def test_load_missing_or_malformed_status_file(svc):
    assert svc._load_image_statuses()==({},False)
    svc._status_file.parent.mkdir(parents=True,exist_ok=True)
    svc._status_file.write_text('{broken json',encoding='utf-8')
    assert svc._load_image_statuses()==({},False)
    svc._status_file.write_text('[]',encoding='utf-8')
    assert svc._load_image_statuses()==({},False)


def test_load_legacy_status_migrates(svc,env,sample):
    payload={'images':{'Sample\\N001\\a.png':{'status':'success','error':'old error','manual_reviewed':'1','result':{'date_reason':'old','date_notes':['abc']}}}}
    svc._status_file.parent.mkdir(parents=True,exist_ok=True)
    svc._status_file.write_text(json.dumps(payload),encoding='utf-8')
    loaded,migrated=svc._load_image_statuses()
    assert migrated and loaded['N001/a.png']['status']==2
    assert loaded['N001/a.png']['manual_reviewed']==1
    assert loaded['N001/a.png']['message']=='old error'


def test_load_rejects_invalid_entries(svc):
    svc._status_file.parent.mkdir(parents=True,exist_ok=True)
    svc._status_file.write_text(json.dumps({'invalid':123,'onepiece':{'status':1},'N/a.png':{'status':9,'result':None,'manual_reviewed':0}}))
    value,migrated=svc._load_image_statuses()
    assert list(value)==['N/a.png'] and value['N/a.png']['status']==0
    assert not migrated


def test_running_status_recovered_on_restart(env,sample):
    path=env.out/'saved.json';path.parent.mkdir(parents=True)
    path.write_text(json.dumps({'N001/a.png':state(env,sample,1)}))
    service=env.m.UnimelbServiceImpl(status_file=path)
    assert service._image_statuses['N001/a.png']['status']==3
    assert 'restarted' in service._image_statuses['N001/a.png']['message'].lower()
    assert json.loads(path.read_text())['N001/a.png']['status']==3


def test_mark_interrupted_only_matching_task(svc,env,sample):
    svc._image_statuses={'N001/a.png':state(env,sample,1),'OTHER/other.png':dict(state(env,sample,1),task_uuid='OTHER')}
    assert svc._mark_interrupted_running_images_failed_locked()
    assert svc._image_statuses['N001/a.png']['status']==3
    assert svc._image_statuses['OTHER/other.png']['status']==1
    assert not svc._mark_interrupted_running_images_failed_locked()


def test_restore_empty_and_all_states(svc,env,sample):
    assert svc._restore_shared_task_locked([]) is None
    assert svc._restore_shared_task_locked([sample]) is None
    for status in (0,1,2,3):
        svc._image_statuses={'N001/a.png':state(env,sample,status,message='fail' if status==3 else None)}
        record=svc._restore_shared_task_locked([sample])
        assert record['status']==status and record['total']==1
        assert len(record['successful'])==(status==2)
        assert len(record['failed'])==(status==3)
        assert len(record['running'])==(status==1)
        assert record['finished_at'] is not None if status in (2,3) else record['finished_at'] is None


def test_restore_task_with_mixed_states(svc,env,sample):
    other=sample.parent/'b.png';other.write_bytes(b'x')
    svc._image_statuses={'N001/a.png':state(env,sample,2,updated_at=3),'N001/b.png':state(env,other,0,updated_at=4)}
    r=svc._restore_shared_task_locked([sample,other])
    assert r['status']==0 and r['completed']==1 and r['success_count']==1


def test_sync_missing_adds_and_skips(svc,env,sample,monkeypatch):
    svc._sync_missing_image_states_locked([sample])
    assert svc._image_statuses['N001/a.png']['status']==0
    svc._sync_missing_image_states_locked([sample])
    assert len(svc._image_statuses)==1
    other=sample.parent/'b.png';other.write_bytes(b'x')
    monkeypatch.setattr(svc,'_persist_image_statuses_locked',Mock(side_effect=OSError('no disk')))
    with pytest.raises(OSError): svc._sync_missing_image_states_locked([other])
    assert 'N001/b.png' not in svc._image_statuses

@pytest.mark.asyncio
async def test_get_image_select_populates_response(svc,env,sample):
    svc._image_statuses['N001/a.png']=state(env,sample,2,manual_reviewed=1,result={'date':'2026','date_conf':0.9})
    r=await svc.get_image_select()
    assert r.total==1 and r.folders=={'SAMPLE':1}
    assert r.status_counts['2']==1
    item=r.images[0]
    assert item.image_key=='N001/a.png' and item.manual_reviewed==1
    assert item.result['date'].value=='2026' and item.result['date'].confidence==0.9

@pytest.mark.asyncio
async def test_get_image_select_adds_missing_statuses(svc,env,sample):
    r=await svc.get_image_select()
    assert r.total==1 and r.images[0].status==0
    assert 'N001/a.png' in svc._image_statuses


def test_format_ocr_results_variants(svc):
    assert svc._format_ocr_result_for_response(None)=={}
    assert svc._format_ocr_result_for_response([])=={}
    assert svc._format_ocr_result_for_response({})=={}
    actual=svc._format_ocr_result_for_response({'date_verbatim':'2020','date_corrected':'2026','date_conf':0.55,'date_tokens':[1,2],'date_tokens_bbox':{'x0':1},'date_notes':['x']})
    assert actual['date']['value']=='2026'
    assert actual['date']['tokens']==[1,2]
    assert actual['date']['tokens_bbox']=={'x0':1}
    original=svc._format_ocr_result_for_response({'date':'exact','date_tokens':'wrong','date_tokens_bbox':[]})
    assert original['date']['value']=='exact' and original['date']['tokens']==[] and original['date']['tokens_bbox'] is None


def test_reason_keys_and_helpers(svc,monkeypatch):
    data={'date_reason':'test','date_notes':['x']}
    normalized,migrated=svc._normalize_result_reason_keys(data)
    assert normalized['date_notes']
    assert not migrated or normalized!=data
    result,migrated=svc._normalize_result_reason_keys({})
    assert result=={} and not migrated
    assert isinstance(svc._note_values('a'),list)
    assert isinstance(svc._split_reason_and_notes('why',['x']),tuple)


def test_reservation_release_only_owner(svc):
    a=object();b=object()
    svc._rerun_image_owners={'N/a':a,'N/b':b}
    svc._release_rerun_image_owners(a,{'N/a','N/b','missing'})
    assert svc._rerun_image_owners=={'N/b':b}


def test_fail_unfinished_filters_by_status_and_task(svc,env,sample):
    svc._image_statuses={'N/a':state(env,sample,1),'N/b':state(env,sample,2),'N/c':state(env,sample,1,task_uuid='another')}
    svc._fail_unfinished_images(env.m.UNIMELB_SHARED_TASK_UUID,'failure',image_keys={'N/a','N/b'})
    assert svc._image_statuses['N/a']['status']==3
    assert svc._image_statuses['N/b']['status']==2
    assert svc._image_statuses['N/c']['status']==1


def test_fail_unfinished_with_ownership(svc,env,sample):
    owned=object();other=object()
    svc._image_statuses={'N/a':state(env,sample,1),'N/b':state(env,sample,1)}
    svc._rerun_image_owners={'N/a':owned,'N/b':other}
    svc._fail_unfinished_images(env.m.UNIMELB_SHARED_TASK_UUID,'stopped',reservation_token=owned)
    assert svc._image_statuses['N/a']['status']==3
    assert svc._image_statuses['N/b']['status']==1
    svc._fail_unfinished_images(env.m.UNIMELB_SHARED_TASK_UUID,'no-op',reservation_token=object())


def test_set_image_status_and_serialized_result(svc,env,sample):
    svc._set_image_status('N001/a.png','task',1,'working')
    assert svc._image_statuses['N001/a.png']['message']=='working'
    token=object();svc._rerun_image_owners={'N001/a.png':token}
    svc._set_serialized_image_result('N001/a.png','task',2,{'date':'2026'},next_image_key='N001/b.png',reservation_token=token)
    assert svc._image_statuses['N001/a.png']['result']=={'date':'2026'}
    assert svc._image_statuses['N001/b.png']['status']==1
    assert 'N001/a.png' not in svc._rerun_image_owners


def test_set_image_result_uses_codec(svc,env,sample,monkeypatch):
    monkeypatch.setattr(env.m,'extract_ocr_result',lambda x:{'architect':'Joe'})
    svc._set_image_result('N001/a.png','task',2,record=object())
    assert svc._image_statuses['N001/a.png']['result']=={'architect':'Joe'}

@pytest.mark.parametrize('bad', ['', '   '])
@pytest.mark.asyncio
async def test_manual_review_rejects_empty_key(svc,bad):
    with pytest.raises(module('services.Unimelb_service').UnimelbTaskRunningError):
        await svc.update_ocr_image_result(bad,{'date':'2026'})

@pytest.mark.asyncio
async def test_manual_review_rejects_empty_result(svc):
    with pytest.raises(module('services.Unimelb_service').UnimelbTaskRunningError):
        await svc.update_ocr_image_result('N001/a.png',{})

@pytest.mark.asyncio
async def test_manual_review_missing_image(svc,sample):
    with pytest.raises(FileNotFoundError):await svc.update_ocr_image_result('N001/missing.png',{'date':'2026'})

@pytest.mark.asyncio
async def test_manual_review_success_and_merge(svc,env,sample):
    svc._image_statuses['N001/a.png']=state(env,sample,2,result={'architect':'old','date_conf':0.9})
    r=await svc.update_ocr_image_result('N001/a.png',{'date':'2026','architect':None})
    assert r.manual_reviewed==1 and r.result=={'date':'2026'}
    final=svc._image_statuses['N001/a.png']
    assert final['result']['architect']=='old' and final['result']['date']=='2026'
    assert final['manual_reviewed']==1 and svc._rerun_image_owners=={}

@pytest.mark.asyncio
async def test_manual_review_creates_initial_state(svc,env,sample):
    r=await svc.update_ocr_image_result('N001/a.png',{'drawing_number':'A-1'})
    assert r.result['drawing_number']=='A-1'
    assert svc._image_statuses['N001/a.png']['status']==0

@pytest.mark.asyncio
async def test_manual_review_running_or_reserved(svc,env,sample):
    exc=module('services.Unimelb_service').UnimelbTaskRunningError
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    with pytest.raises(exc):await svc.update_ocr_image_result('N001/a.png',{'date':'2026'})
    svc._image_statuses['N001/a.png']['status']=2
    svc._rerun_image_owners['N001/a.png']=object()
    with pytest.raises(exc):await svc.update_ocr_image_result('N001/a.png',{'date':'2026'})

@pytest.mark.asyncio
async def test_manual_review_save_failure_rolls_back(svc,env,sample,monkeypatch):
    original=state(env,sample,2,result={'date':'old'})
    svc._image_statuses['N001/a.png']=original.copy()
    monkeypatch.setattr(svc,'_persist_image_statuses_locked',Mock(side_effect=OSError('disk full')))
    with pytest.raises(OSError):await svc.update_ocr_image_result('N001/a.png',{'date':'new'})
    assert svc._image_statuses['N001/a.png']['result']['date']=='old'
    assert svc._rerun_image_owners=={}

@pytest.mark.asyncio
async def test_create_task_no_images(svc):
    with pytest.raises(FileNotFoundError):await svc.create_ocr_task(False)

@pytest.mark.asyncio
async def test_rerun_missing_images_and_bad_key(svc,sample):
    with pytest.raises(FileNotFoundError):await svc.rerun_ocr_images(['N001/missing.png'])
    sample.unlink()
    with pytest.raises(FileNotFoundError):await svc.rerun_ocr_images(['N001/a.png'])

@pytest.mark.asyncio
async def test_create_task_background_job_without_process_pool(svc,env,sample,monkeypatch):
    outcome=env.m.OcrImageOutcome(ok=True,result={'date':'2026'},summary_row={'image':'a'})
    monkeypatch.setattr(svc,'_submit_ocr_image',AsyncMock(return_value=outcome))
    mock_summary=AsyncMock(return_value='summary.xlsx')
    monkeypatch.setattr(svc,'_submit_batch_summary',mock_summary)
    r=await svc.create_ocr_task(False)
    assert r.started and r.total==1 and r.status==1
    await asyncio.gather(*list(svc._background_tasks))
    assert svc._image_statuses['N001/a.png']['status']==2
    assert svc._image_statuses['N001/a.png']['result']=={'date':'2026'}
    mock_summary.assert_awaited_once()
    again=await svc.create_ocr_task(False)
    assert not again.started

@pytest.mark.asyncio
async def test_create_task_force_reruns_completed(svc,env,sample,monkeypatch):
    monkeypatch.setattr(svc,'_submit_ocr_image',AsyncMock(return_value=env.m.OcrImageOutcome(ok=False,error='failed')))
    initial=await svc.create_ocr_task(True)
    assert initial.started
    await asyncio.gather(*list(svc._background_tasks))
    assert svc._image_statuses['N001/a.png']['status']==3
    again=await svc.create_ocr_task(True)
    assert again.started
    await asyncio.gather(*list(svc._background_tasks))

@pytest.mark.asyncio
async def test_create_task_reservation_conflict(svc,sample):
    svc._rerun_image_owners['N001/a.png']=object()
    with pytest.raises(module('services.Unimelb_service').UnimelbTaskRunningError):await svc.create_ocr_task(True)

@pytest.mark.asyncio
async def test_rerun_background_task(svc,env,sample,monkeypatch):
    monkeypatch.setattr(svc,'_submit_ocr_image',AsyncMock(return_value=env.m.OcrImageOutcome(ok=True,result={'date':'ok'})))
    r=await svc.rerun_ocr_images(['N001/a.png'])
    assert r.selected_total==1 and r.image_keys==['N001/a.png']
    await asyncio.gather(*list(svc._background_tasks))
    assert svc._image_statuses['N001/a.png']['status']==2

@pytest.mark.asyncio
async def test_rerun_existing_owner_conflict(svc,sample):
    svc._rerun_image_owners['N001/a.png']=object()
    with pytest.raises(module('services.Unimelb_service').UnimelbTaskRunningError):await svc.rerun_ocr_images(['N001/a.png'])

@pytest.mark.asyncio
async def test_rerun_failure_to_persist_releases_owners(svc,sample,monkeypatch):
    monkeypatch.setattr(svc,'_persist_image_statuses_locked',Mock(side_effect=OSError('disk')))
    with pytest.raises(OSError):await svc.rerun_ocr_images(['N001/a.png'])
    assert svc._rerun_image_owners=={}


def test_build_excel_reviewed_unreviewed_and_low_conf(env,monkeypatch):
    cfg=SimpleNamespace(conf=SimpleNamespace(t_high=0.8))
    monkeypatch.setattr(env.m.Config,'load',staticmethod(lambda path:cfg))
    cls=env.m.UnimelbServiceImpl
    states={'N/a.png':{'manual_reviewed':0,'result':{'date':'2026','date_conf':0.5}},'N/b.png':{'manual_reviewed':1,'result':{'date':'2025','date_conf':0.99}}}
    binary=cls._build_results_excel(states)
    sheet=load_workbook(BytesIO(binary)).active
    assert sheet['A1'].value=='image_key'
    headers=[cell.value for cell in sheet[1]]
    assert 'date' in headers and 'date_conf' in headers
    c=headers.index('date')+1
    assert sheet.cell(2,c).value=='2026'
    assert sheet.cell(2,c+1).value==0.5
    assert sheet.cell(2,c).fill.fgColor.rgb=='00FFC7CE'
    assert sheet.cell(3,c+1).value is None
    assert sheet.cell(3,1).fill.fgColor.rgb=='00C6EFCE'


def test_build_excel_all_reviewed_omits_confidence(env,monkeypatch):
    monkeypatch.setattr(env.m.Config,'load',staticmethod(lambda path:SimpleNamespace(conf=SimpleNamespace(t_high=0.8))))
    xls=env.m.UnimelbServiceImpl._build_results_excel({'N/a':{'manual_reviewed':1,'result':{'scale':'1:100'}}})
    sheet=load_workbook(BytesIO(xls)).active
    headers=[c.value for c in sheet[1]]
    assert 'scale' in headers and 'scale_conf' not in headers


def test_export_results_success_filters(svc,env,sample,monkeypatch):
    monkeypatch.setattr(env.m.Config,'load',staticmethod(lambda path:SimpleNamespace(conf=SimpleNamespace(t_high=0.8))))
    svc._image_statuses['N001/a.png']=state(env,sample,2,result={'date':'2026'})
    xls=svc._export_ocr_results(None,status=2,folder='SAMPLE')
    sheet=load_workbook(BytesIO(xls)).active
    assert sheet.max_row==2 and sheet['A2'].value=='N001/a.png'
    empty=svc._export_ocr_results(None,status=3)
    assert load_workbook(BytesIO(empty)).active.max_row==1


def test_export_rejects_busy_status_and_missing_keys(svc,env,sample):
    exc=module('services.Unimelb_service').UnimelbTaskRunningError
    with pytest.raises(exc):svc._export_ocr_results(None,status=1)
    with pytest.raises(FileNotFoundError):svc._export_ocr_results(['N001/absent.png'])
    svc._image_statuses['N001/a.png']=state(env,sample,2)
    svc._rerun_image_owners['N001/a.png']=object()
    with pytest.raises(exc):svc._export_ocr_results(['N001/a.png'])


def test_export_rejects_in_progress_image(svc,env,sample):
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    with pytest.raises(module('services.Unimelb_service').UnimelbTaskRunningError):svc._export_ocr_results(['N001/a.png'])
    assert svc._rerun_image_owners=={}


def test_export_rejects_missing_file(svc,env,sample):
    svc._image_statuses['N001/a.png']=state(env,sample,2)
    sample.unlink()
    with pytest.raises(FileNotFoundError):svc._export_ocr_results(['N001/a.png'])

@pytest.mark.asyncio
async def test_export_async_wrapper(svc,env,sample,monkeypatch):
    monkeypatch.setattr(svc,'_export_ocr_results',Mock(return_value=b'XLSX'))
    assert await svc.export_ocr_results(['N001/a.png'],status=2,folder='SAMPLE')==b'XLSX'
    svc._export_ocr_results.assert_called_once_with(['N001/a.png'],status=2,folder='SAMPLE')


def test_executor_lifecycle_without_workers(svc,env,monkeypatch):
    fake=Mock()
    constructor=Mock(return_value=fake)
    monkeypatch.setattr(env.m,'ProcessPoolExecutor',constructor)
    executor=svc._get_ocr_executor()
    assert executor is fake and svc._get_ocr_executor() is executor
    constructor.assert_called_once()
    svc._discard_broken_executor(Mock())
    assert svc._ocr_executor is fake
    svc._discard_broken_executor(executor)
    assert svc._ocr_executor is None
    fake.shutdown.assert_called_once_with(wait=False,cancel_futures=True)
    svc._closing=True
    with pytest.raises(RuntimeError):svc._get_ocr_executor()

@pytest.mark.asyncio
async def test_start_close_fake_executor(svc,env,monkeypatch):
    fake=Mock();monkeypatch.setattr(env.m,'ProcessPoolExecutor',Mock(return_value=fake))
    monkeypatch.setattr(svc,'_kill_ocr_executor',lambda executor:executor.shutdown(wait=True))
    await svc.start()
    assert svc._ocr_executor is fake
    await svc.close()
    assert svc._closing and svc._ocr_executor is None
    fake.shutdown.assert_called_once_with(wait=True)
    await svc.close()


def test_kill_executor_fallback_and_new_api(svc):
    killer=Mock();executor=SimpleNamespace(kill_workers=killer)
    svc._kill_ocr_executor(executor)
    killer.assert_called_once()
    proc=Mock();proc.is_alive.return_value=True
    old=SimpleNamespace(_processes={1:proc},shutdown=Mock())
    svc._kill_ocr_executor(old)
    proc.kill.assert_called_once()
    old.shutdown.assert_called_once_with(wait=True,cancel_futures=True)
    dead=Mock();dead.is_alive.return_value=False
    old=SimpleNamespace(_processes={1:dead},shutdown=Mock())
    svc._kill_ocr_executor(old)
    dead.kill.assert_not_called()

@pytest.mark.asyncio
async def test_run_ocr_task_process_failure_paths(svc,env,sample,monkeypatch):
    """Check both infrastructure exceptions and regular worker exceptions."""
    for error in (BrokenProcessPool('worker died'), ValueError('bad OCR')):
        svc._image_statuses={'N001/a.png':state(env,sample,1)}
        token=object();svc._rerun_image_owners={'N001/a.png':token}
        monkeypatch.setattr(svc,'_submit_ocr_image',AsyncMock(side_effect=error))
        await svc._run_ocr_task(env.m.UNIMELB_SHARED_TASK_UUID,[sample],env.out,reservation_token=token)
        assert svc._image_statuses['N001/a.png']['status']==3
        assert type(error).__name__ in svc._image_statuses['N001/a.png']['message'] or 'worker' in svc._image_statuses['N001/a.png']['message'].lower()
        assert not svc._rerun_image_owners

@pytest.mark.asyncio
async def test_run_ocr_task_cancelled(svc,env,sample,monkeypatch):
    token=object();svc._rerun_image_owners['N001/a.png']=token
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    monkeypatch.setattr(svc,'_submit_ocr_image',AsyncMock(side_effect=asyncio.CancelledError()))
    with pytest.raises(asyncio.CancelledError):
        await svc._run_ocr_task(env.m.UNIMELB_SHARED_TASK_UUID,[sample],env.out,reservation_token=token)
    assert svc._image_statuses['N001/a.png']['status']==3
    assert not svc._rerun_image_owners

@pytest.mark.asyncio
async def test_run_ocr_task_cancellation_after_saving_result(svc,env,sample,monkeypatch):
    token=object();svc._rerun_image_owners['N001/a.png']=token
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    monkeypatch.setattr(svc,'_submit_ocr_image',AsyncMock(return_value=env.m.OcrImageOutcome(ok=True,result={'date':'ok'})))
    monkeypatch.setattr(svc,'_await_shielded_task',AsyncMock(return_value=(env.m.OcrImageOutcome(ok=True,result={'date':'ok'}),asyncio.CancelledError())))
    with pytest.raises(asyncio.CancelledError):
        await svc._run_ocr_task(env.m.UNIMELB_SHARED_TASK_UUID,[sample],env.out,reservation_token=token)
    assert svc._image_statuses['N001/a.png']['status']==2

@pytest.mark.asyncio
async def test_run_ocr_task_summary_failure_and_cancel(svc,env,sample,monkeypatch):
    outcome=env.m.OcrImageOutcome(ok=True,result={'date':'ok'},summary_row={'image':'a'})
    monkeypatch.setattr(svc,'_submit_ocr_image',AsyncMock(return_value=outcome))
    monkeypatch.setattr(svc,'_submit_batch_summary',AsyncMock(side_effect=ValueError('summary failed')))
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    await svc._run_ocr_task(env.m.UNIMELB_SHARED_TASK_UUID,[sample],env.out,reservation_token=object())
    # The result was saved even though summary creation failed.
    assert svc._image_statuses['N001/a.png']['result']=={'date':'ok'}

@pytest.mark.asyncio
async def test_await_shielded_task_deferred_cancel(svc):
    proceed=asyncio.Event()
    async def delayed():
        await proceed.wait()
        return 'finished'
    task=asyncio.create_task(delayed())
    outer=asyncio.create_task(svc._await_shielded_task(task))
    await asyncio.sleep(0)
    outer.cancel()
    await asyncio.sleep(0)
    proceed.set()
    result,cancellation=await outer
    assert result=='finished' and isinstance(cancellation,asyncio.CancelledError)

@pytest.mark.asyncio
async def test_submit_image_and_summary_with_thread_executor(svc,env,monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    job=env.m.OcrImageJob(image_path='fake.png',node_id='N',item_output_dir='out')
    with ThreadPoolExecutor(max_workers=1) as pool:
        monkeypatch.setattr(svc,'_get_ocr_executor',lambda:pool)
        monkeypatch.setattr(env.m,'process_ocr_image',lambda task: env.m.OcrImageOutcome(ok=True,result={'date':task.node_id}))
        monkeypatch.setattr(env.m,'write_worker_batch_summary',lambda output,rows:'summary.xlsx')
        result=await svc._submit_ocr_image(job)
        summary=await svc._submit_batch_summary(env.out,[{'image':'a'}])
    assert result.ok and result.result['date']=='N'
    assert summary=='summary.xlsx'

@pytest.mark.asyncio
async def test_submit_image_discards_broken_pool(svc,env,monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    job=env.m.OcrImageJob(image_path='fake.png',node_id='N',item_output_dir='out')
    discarded=Mock();monkeypatch.setattr(svc,'_discard_broken_executor',discarded)
    def broken(*args): raise BrokenProcessPool('pool died')
    with ThreadPoolExecutor(max_workers=1) as pool:
        monkeypatch.setattr(svc,'_get_ocr_executor',lambda:pool)
        monkeypatch.setattr(env.m,'process_ocr_image',broken)
        with pytest.raises(BrokenProcessPool):await svc._submit_ocr_image(job)
        discarded.assert_called_once_with(pool)

@pytest.mark.asyncio
async def test_submit_summary_discards_broken_pool(svc,env,monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    discarded=Mock();monkeypatch.setattr(svc,'_discard_broken_executor',discarded)
    def broken(*args): raise BrokenProcessPool('pool died')
    with ThreadPoolExecutor(max_workers=1) as pool:
        monkeypatch.setattr(svc,'_get_ocr_executor',lambda:pool)
        monkeypatch.setattr(env.m,'write_worker_batch_summary',broken)
        with pytest.raises(BrokenProcessPool):await svc._submit_batch_summary(env.out,[])
        discarded.assert_called_once_with(pool)


def test_result_codec_forwarding_methods(svc,env,monkeypatch):
    marker={'x0':1}
    monkeypatch.setattr(env.m,'evidence_regions_bbox',lambda signals: marker)
    monkeypatch.setattr(env.m,'tokens_bbox',lambda tokens,token_map: marker)
    monkeypatch.setattr(env.m,'extract_ocr_result',lambda record: {'scale':'1:100'})
    assert svc._evidence_regions_bbox({}) == marker
    assert svc._tokens_bbox([1],{}) == marker
    assert svc._extract_ocr_result(object())=={'scale':'1:100'}


def test_normalize_reason_keys_changes_legacy_fields(svc,env,monkeypatch):
    monkeypatch.setattr(env.m,'split_reason_and_notes',lambda reason,notes:('not_present',['readable']))
    result,migrated=svc._normalize_result_reason_keys({'date_reason':{'old':'format'}})
    assert migrated
    assert result['date_reason']=='not_present' and result['date_notes']==['readable']
    original,migrated=svc._normalize_result_reason_keys({'date_reason':'not_present','date_notes':['readable']})
    assert not migrated and original['date_reason']=='not_present'


def test_display_path_external_fallback(svc,env):
    assert svc._display_image_path(Path('/outside-root/a.png')) == str(Path('/outside-root/a.png'))


def test_save_failure_propagates(svc,env,monkeypatch):
    original=Path.write_text
    def reject(path,*args,**kwargs):
        if str(path).endswith('test-status.json.tmp'):
            raise OSError('disk full')
        return original(path,*args,**kwargs)
    monkeypatch.setattr(Path,'write_text',reject)
    with pytest.raises(OSError):svc._persist_image_statuses_locked()


def test_kill_executor_ignores_process_not_found(svc):
    proc=Mock();proc.is_alive.return_value=True;proc.kill.side_effect=ProcessLookupError()
    executor=SimpleNamespace(_processes={1:proc},shutdown=Mock())
    svc._kill_ocr_executor(executor)
    executor.shutdown.assert_called_once()


def test_build_excel_high_confidence_not_highlighted(env,monkeypatch):
    monkeypatch.setattr(env.m.Config,'load',staticmethod(lambda path:SimpleNamespace(conf=SimpleNamespace(t_high=0.8))))
    result=env.m.UnimelbServiceImpl._build_results_excel({'N/a.png':{'manual_reviewed':0,'result':{'date':'2026','date_conf':0.95}}})
    sheet=load_workbook(BytesIO(result)).active
    col=[cell.value for cell in sheet[1]].index('date')+1
    assert sheet.cell(2,col).fill.patternType is None

@pytest.mark.asyncio
async def test_manual_review_new_image_save_failure_rolls_back(svc,env,sample,monkeypatch):
    monkeypatch.setattr(svc,'_persist_image_statuses_locked',Mock(side_effect=OSError('no space')))
    with pytest.raises(OSError):await svc.update_ocr_image_result('N001/a.png',{'date':'2026'})
    assert 'N001/a.png' not in svc._image_statuses
    assert not svc._rerun_image_owners

@pytest.mark.asyncio
async def test_rerun_running_state_conflict(svc,env,sample):
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    with pytest.raises(module('services.Unimelb_service').UnimelbTaskRunningError):
        await svc.rerun_ocr_images(['N001/a.png'])
    assert not svc._rerun_image_owners

@pytest.mark.asyncio
async def test_create_task_running_state_force_conflict(svc,env,sample):
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    with pytest.raises(module('services.Unimelb_service').UnimelbTaskRunningError):
        await svc.create_ocr_task(True)
    assert not svc._rerun_image_owners

@pytest.mark.asyncio
async def test_create_task_persistence_failure_rolls_back(svc,env,sample,monkeypatch):
    monkeypatch.setattr(svc,'_persist_image_statuses_locked',Mock(side_effect=OSError('storage offline')))
    with pytest.raises(OSError):await svc.create_ocr_task(False)
    assert not svc._rerun_image_owners and svc._image_statuses=={}

@pytest.mark.asyncio
async def test_create_task_cannot_schedule_job_rolls_back(svc,env,sample,monkeypatch):
    def fail_to_schedule(coro):
        coro.close()
        raise RuntimeError('no scheduler')
    monkeypatch.setattr(env.m.asyncio,'create_task',fail_to_schedule)
    with pytest.raises(RuntimeError):await svc.create_ocr_task(False)
    assert not svc._rerun_image_owners
    assert svc._image_statuses['N001/a.png']['status']==3

@pytest.mark.asyncio
async def test_rerun_cannot_schedule_job_rolls_back(svc,env,sample,monkeypatch):
    def fail_to_schedule(coro):
        coro.close()
        raise RuntimeError('no scheduler')
    monkeypatch.setattr(env.m.asyncio,'create_task',fail_to_schedule)
    with pytest.raises(RuntimeError):await svc.rerun_ocr_images(['N001/a.png'])
    assert not svc._rerun_image_owners
    assert svc._image_statuses['N001/a.png']['status']==3

@pytest.mark.asyncio
async def test_create_task_skips_running_image_owned_by_existing_task(svc,env,sample):
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    result=await svc.create_ocr_task(False)
    assert not result.started and result.status==1

@pytest.mark.asyncio
async def test_create_task_skips_images_claimed_by_another_request(svc,env,sample):
    svc._rerun_image_owners['N001/a.png']=object()
    result=await svc.create_ocr_task(False)
    assert not result.started
    assert 'N001/a.png' in svc._rerun_image_owners

@pytest.mark.asyncio
async def test_run_ocr_task_cancellation_after_summary(svc,env,sample,monkeypatch):
    outcome=env.m.OcrImageOutcome(ok=True,result={'date':'2026'},summary_row={'image':'a'})
    monkeypatch.setattr(svc,'_submit_ocr_image',AsyncMock(return_value=outcome))
    monkeypatch.setattr(svc,'_submit_batch_summary',AsyncMock(return_value='summary'))
    monkeypatch.setattr(svc,'_await_shielded_task',AsyncMock(side_effect=[(outcome,None),('summary',asyncio.CancelledError())]))
    svc._image_statuses['N001/a.png']=state(env,sample,1)
    with pytest.raises(asyncio.CancelledError):
        await svc._run_ocr_task(env.m.UNIMELB_SHARED_TASK_UUID,[sample],env.out,reservation_token=object())
    assert svc._image_statuses['N001/a.png']['status']==2

@pytest.mark.asyncio
async def test_close_cancels_background_tasks(svc):
    async def never_finishes():
        await asyncio.Event().wait()
    task=asyncio.create_task(never_finishes())
    svc._background_tasks.add(task)
    await asyncio.sleep(0)
    await svc.close()
    assert task.cancelled() and svc._closing

@pytest.mark.asyncio
async def test_close_cleans_tasks_when_executor_kill_fails(svc,monkeypatch):
    async def never_finishes():
        await asyncio.Event().wait()
    task=asyncio.create_task(never_finishes())
    svc._background_tasks.add(task)
    fake=Mock();svc._ocr_executor=fake
    monkeypatch.setattr(svc,'_kill_ocr_executor',Mock(side_effect=RuntimeError('shutdown failed')))
    await asyncio.sleep(0)
    with pytest.raises(RuntimeError):await svc.close()
    assert task.cancelled() and svc._ocr_executor is None
