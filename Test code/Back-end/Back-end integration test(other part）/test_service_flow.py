from io import BytesIO
from pathlib import Path
import asyncio
from concurrent.futures.process import BrokenProcessPool
from copy import deepcopy
import json
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient

from openpyxl import load_workbook
from conftest import REVIEW_VALUES, drain
from uma.web.app import create_app
from uma.web.services import UnimelbTaskRunningError
from uma.web.services.impl import Unimelb_service_impl as service_module
from uma.web.workers.Unimelb_ocr_worker import OcrImageOutcome

def workbook_rows(content):
    book = load_workbook(BytesIO(content))
    try:
        sheet = book.active
        headers = [cell.value for cell in sheet[1]]
        return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    finally:
        book.close()


def seed_completed(service):
    for image in service._collect_images():
        service._set_serialized_image_result(
            service._image_key(image), service_module.UNIMELB_SHARED_TASK_UUID, 2,
            {"drawing_number": "003/A1", "drawing_number_conf": .5,
             "drawing_number_verbatim": "003/A1", "drawing_number_tokens": [1],
             "drawing_number_tokens_bbox": {"x0": 1, "y0": 2, "x1": 20, "y1": 30}},
        )


def test_real_service_http_listing_review_reload_and_excel(service_env):
    root, paths, make = service_env
    service = make()
    with TestClient(create_app(unimel_service=service)) as client:
        listing = client.get("/Unimelb/ImageSelect").json()["result"]
        assert listing["total"] == 3 and listing["folders"] == {"OTHER": 1, "YFA": 2}
        assert listing["status_counts"]["0"] == 3
        image = next(row for row in listing["images"] if row["node_id"] == "100")
        assert client.get(image["image_url"]).content == paths[0].read_bytes()
        seed_completed(service)
        service._image_statuses["101/drawing.png"]["result"]["drawing_number_conf"] = .95
        saved = client.post("/Unimelb/ocr-results/manual-review", json={
            "image_key": image["image_key"], "result": REVIEW_VALUES,
        })
        assert saved.status_code == 200 and saved.json()["result"]["manual_reviewed"] == 1
        reread = client.get("/Unimelb/ImageSelect").json()["result"]["images"]
        updated = next(row for row in reread if row["node_id"] == "100")
        assert updated["result"]["drawing_number"]["value"] == "003/A2"
        downloaded = client.post("/Unimelb/ocr-results/export", json={"image_keys": [image["image_key"]]})
        assert downloaded.status_code == 200
        rows = workbook_rows(downloaded.content)
        assert rows[0]["drawing_number"] == "003/A2"
        assert "drawing_number_conf" not in rows[0]

        # Read the real mixed-review workbook returned by the HTTP endpoint.
        mixed = client.post("/Unimelb/ocr-results/export", json={})
        assert mixed.status_code == 200
        book = load_workbook(BytesIO(mixed.content))
        try:
            sheet = book.active
            headers = [cell.value for cell in sheet[1]]
            assert headers[0] == "image_key"
            assert "project_name" in headers and "project_name_conf" in headers
            number_column = headers.index("drawing_number") + 1
            confidence_column = headers.index("drawing_number_conf") + 1
            row_by_key = {sheet.cell(row, 1).value: row for row in range(2, sheet.max_row + 1)}
            assert list(row_by_key) == ["102/drawing.png", "100/drawing.png", "101/drawing.png"]
            for key, expected_value, confidence in [
                ("102/drawing.png", "003/A1", .5),
                ("100/drawing.png", "003/A2", None),
                ("101/drawing.png", "003/A1", .95),
            ]:
                row = row_by_key[key]
                value_cell = sheet.cell(row, number_column)
                assert value_cell.value == expected_value
                assert value_cell.data_type == "s" and value_cell.number_format == "@"
                assert sheet.cell(row, confidence_column).value == confidence
            reviewed = row_by_key["100/drawing.png"]
            assert sheet.cell(reviewed, 1).fill.fgColor.rgb.endswith("C6EFCE")
            low = row_by_key["102/drawing.png"]
            assert sheet.cell(low, number_column).fill.fgColor.rgb.endswith("FFC7CE")
            assert sheet.cell(low, confidence_column).fill.fgColor.rgb.endswith("FFC7CE")
            high = row_by_key["101/drawing.png"]
            assert sheet.cell(high, number_column).fill.patternType is None
            assert sheet.freeze_panes == "B2" and sheet.auto_filter.ref == sheet.dimensions
        finally:
            book.close()
    restored = make()
    restored_state = restored._image_statuses[image["image_key"]]
    assert restored_state["result"]["drawing_number"] == "003/A2"
    assert restored_state["manual_reviewed"] == 1 and restored_state["status"] == 2
    assert restored_state["result"]["project_name"] == REVIEW_VALUES["project_name"]
    assert not service._rerun_image_owners


@pytest.mark.parametrize("filters,expected", [
    ({}, ["102/drawing.png", "100/drawing.png", "101/drawing.png"]),
    ({"folder": "yfa"}, ["100/drawing.png", "101/drawing.png"]),
    ({"folder": "OTHER", "status": 2}, ["102/drawing.png"]),
    ({"status": 0}, []),
    ({"image_keys": ["101/drawing.png", "101/drawing.png"]}, ["101/drawing.png"]),
])
def test_real_export_filters_sorting_and_empty_workbook(service_env, filters, expected):
    _, _, make = service_env
    service = make()
    seed_completed(service)
    with TestClient(create_app(unimel_service=service)) as client:
        response = client.post("/Unimelb/ocr-results/export", json=filters)
        assert response.status_code == 200
        assert [row["image_key"] for row in workbook_rows(response.content)] == expected
    assert not service._rerun_image_owners


@pytest.mark.parametrize("mode,code", [("unknown", 404), ("reserved", 409), ("running", 409), ("status-filter", 409)])
def test_export_rejects_missing_and_busy_images(service_env, mode, code):
    _, _, make = service_env
    service = make()
    seed_completed(service)
    key = "100/drawing.png"
    payload = {"image_keys": [key]}
    if mode == "unknown":
        payload["image_keys"] = ["missing/no.png"]
    elif mode == "reserved":
        service._rerun_image_owners[key] = object()
    elif mode == "running":
        service._image_statuses[key]["status"] = 1
    elif mode == "status-filter":
        payload["status"] = 1
    with TestClient(create_app(unimel_service=service)) as client:
        response = client.post("/Unimelb/ocr-results/export", json=payload)
        assert response.status_code == code and response.json()["success"] is False


def test_create_complete_rerun_and_force_reset_without_ocr(service_env):
    _, _, make = service_env
    service = make()

    async def workflow():
        await service.start()
        first = await service.create_ocr_task(False)
        assert first.started and first.status == 1 and first.total == 3
        await drain(service)
        assert all(state["status"] == 2 for state in service._image_statuses.values())
        assert not (await service.create_ocr_task(False)).started
        await service.update_ocr_image_result("100/drawing.png", REVIEW_VALUES)
        assert service._image_statuses["100/drawing.png"]["manual_reviewed"] == 1
        selected = await service.rerun_ocr_images(["100/drawing.png"])
        assert selected.started and selected.selected_total == 1
        await drain(service)
        assert service._image_statuses["100/drawing.png"]["manual_reviewed"] == 0
        assert (await service.create_ocr_task(True)).started
        await drain(service)
        assert not service._rerun_image_owners
        # The public create response refreshes the cached shared-task summary.
        final = await service.create_ocr_task(False)
        assert not final.started and final.status == 2
        task = service._tasks[service_module.UNIMELB_SHARED_TASK_UUID]
        assert task["status"] == 2 and task["success_count"] == 3
        assert Path(task["output_dir"]).name == service_module.UNIMELB_SHARED_TASK_UUID
        await service.close()
        await service.close()
        with pytest.raises(RuntimeError, match="shutting down"):
            service._get_ocr_executor()

    asyncio.run(workflow())


@pytest.mark.parametrize("operation,busy", [
    ("create", "owner"), ("create", "state"),
    ("rerun", "owner"), ("review", "state"),
])
def test_busy_image_conflicts_leave_existing_state_unchanged(service_env, operation, busy):
    _, _, make = service_env
    service = make()
    seed_completed(service)
    key = "100/drawing.png"
    if busy == "owner":
        service._rerun_image_owners[key] = object()
    else:
        service._image_statuses[key]["status"] = 1
    original = deepcopy(service._image_statuses[key])

    async def workflow():
        with pytest.raises(UnimelbTaskRunningError):
            if operation == "create":
                await service.create_ocr_task(True)
            elif operation == "rerun":
                await service.rerun_ocr_images([key])
            else:
                await service.update_ocr_image_result(key, REVIEW_VALUES)
        assert service._image_statuses[key] == original

    asyncio.run(workflow())


def test_missing_image_root_returns_an_http_error(service_env, monkeypatch):
    root, _, make = service_env
    service = make()
    monkeypatch.setattr(service_module, "PLAN_SAMPLE_DIR", root / "not-there")

    with TestClient(create_app(unimel_service=service)) as client:
        response = client.get("/Unimelb/ImageSelect")
        assert response.status_code == 404 and response.json()["success"] is False
    assert service._collect_images() == []


@pytest.mark.parametrize("operation", ["create", "rerun"])
def test_no_images_cannot_start_a_task(service_env, monkeypatch, operation):
    root, _, make = service_env
    service = make()
    empty = root / "empty"
    empty.mkdir()
    monkeypatch.setattr(service_module, "PLAN_SAMPLE_DIR", empty)
    with pytest.raises(FileNotFoundError, match="No images"):
        asyncio.run(service.create_ocr_task(False) if operation == "create" else service.rerun_ocr_images([]))


def test_review_missing_image_does_not_change_persisted_data(service_env):
    _, _, make = service_env
    service = make()
    with TestClient(create_app(unimel_service=service)) as client:
        response = client.post("/Unimelb/ocr-results/manual-review", json={
            "image_key": "missing/file.png", "result": REVIEW_VALUES,
        })
        assert response.status_code == 404
    assert service._image_statuses == {} and not service._rerun_image_owners


def test_unknown_rerun_and_path_traversal_are_rejected(service_env):
    root, paths, make = service_env
    service = make()
    with pytest.raises(FileNotFoundError):
        asyncio.run(service.rerun_ocr_images(["missing/no.png"]))
    assert service.resolve_image("yfa", "100", "drawing.png") == paths[0]
    for folder, node, name in [("missing", "100", "drawing.png"), ("YFA", "../100", "drawing.png"),
                               ("YFA", "100", "../drawing.png"), ("YFA", "100", "notes.txt")]:
        assert service.resolve_image(folder, node, name) is None
    outside = root / "outside.png"
    outside.write_bytes(b"outside")
    (paths[0].parent / "escape.png").symlink_to(outside)
    assert service.resolve_image("YFA", "100", "escape.png") is None


@pytest.mark.parametrize("operation", ["create", "review"])
def test_disk_failure_rolls_back_state_and_reservations(service_env, monkeypatch, operation):
    _, _, make = service_env
    service = make()
    seed_completed(service)
    if operation == "create":
        service._image_statuses.clear()
    original = deepcopy(service._image_statuses)
    with monkeypatch.context() as patch:
        patch.setattr(Path, "replace", Mock(side_effect=OSError("disk full")))
        with pytest.raises(OSError, match="disk full"):
            if operation == "create":
                asyncio.run(service.create_ocr_task(True))
            else:
                asyncio.run(service.update_ocr_image_result("100/drawing.png", REVIEW_VALUES))
    assert service._image_statuses == original and not service._rerun_image_owners


@pytest.mark.parametrize("payload", ["invalid-json", {"bad": 2, "single": {}, "100/drawing.png": {
    "status": "success", "manual_reviewed": "1", "result": {"drawing_number_reason": ["ambiguous", "note"]},
}}, {"images": {"YFA\\100\\drawing.png": {
    "status": "running", "manual_reviewed": "0", "task_uuid": service_module.UNIMELB_SHARED_TASK_UUID,
    "result": {"drawing_number_notes": "legacy note"},
}}}])
def test_legacy_and_broken_status_files_reload_safely(service_env, payload):
    root, _, make = service_env
    path = root / "output/status.json"
    path.parent.mkdir(parents=True)
    path.write_text(payload if isinstance(payload, str) else json.dumps(payload))
    service = make()
    listing = asyncio.run(service.get_image_select())
    assert listing.total == 3
    assert all(row.status in {0, 1, 2, 3} for row in listing.images)
    assert all(row.manual_reviewed in {0, 1} for row in listing.images)
    if isinstance(payload, dict) and "images" in payload:
        assert service._image_statuses["100/drawing.png"]["status"] == 3
        assert "restarted" in service._image_statuses["100/drawing.png"]["message"]
    assert isinstance(json.loads(path.read_text()), dict)


@pytest.mark.parametrize("error", [ValueError("bad worker result"), BrokenProcessPool("worker stopped")])
def test_worker_dispatch_failure_marks_unfinished_images_failed(service_env, monkeypatch, error):
    _, _, make = service_env
    service = make()
    monkeypatch.setattr(service_module, "process_ocr_image", Mock(side_effect=error))

    async def workflow():
        assert (await service.create_ocr_task(False)).started
        await drain(service)
        assert all(state["status"] == 3 for state in service._image_statuses.values())
        assert not service._rerun_image_owners
        if isinstance(error, BrokenProcessPool):
            assert service._ocr_executor is None
        await service.close()
    asyncio.run(workflow())


def test_failed_outcome_is_persisted_without_stopping_next_image(service_env, monkeypatch):
    _, _, make = service_env
    service = make()
    original = service_module.process_ocr_image
    monkeypatch.setattr(service_module, "process_ocr_image", lambda job: (
        OcrImageOutcome(ok=False, error="unreadable") if job.node_id == "100" else original(job)
    ))

    async def workflow():
        await service.create_ocr_task(False)
        await drain(service)
        assert service._image_statuses["100/drawing.png"]["status"] == 3
        assert service._image_statuses["101/drawing.png"]["status"] == 2
        task = service._restore_shared_task_locked(service._collect_images())
        assert task["status"] == 3 and task["failed_count"] == 1
        await service.close()
    asyncio.run(workflow())


def test_cancelled_batch_keeps_finished_result_and_releases_remaining_images(service_env, monkeypatch):
    _, _, make = service_env
    service = make()

    async def workflow():
        entered, release = asyncio.Event(), asyncio.Event()
        async def controlled_inference(job):
            entered.set()
            await release.wait()
            return OcrImageOutcome(ok=True, result={"drawing_number": "S24"})
        monkeypatch.setattr(service, "_submit_ocr_image", controlled_inference)
        await service.create_ocr_task(False)
        await asyncio.wait_for(entered.wait(), timeout=2)
        task = next(iter(service._background_tasks))
        task.cancel()
        await asyncio.sleep(0)
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await asyncio.wait_for(task, timeout=2)
        assert service._image_statuses["102/drawing.png"]["status"] == 2
        assert service._image_statuses["100/drawing.png"]["status"] == 3
        assert not service._rerun_image_owners
        await service.close()
    asyncio.run(workflow())
